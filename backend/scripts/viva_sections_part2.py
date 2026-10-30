"""
Part 2 of EventIQ Viva Preparation Guide: Sections 11 to 20
Exhaustive Academic & Practical Explanations (Fully Expanded for Viva Mastery)
"""
from docx.shared import Pt, RGBColor
from viva_guide_builder import (
    add_heading_1, add_heading_2, add_heading_3,
    add_p, add_bullet, add_callout, add_formula_box,
    add_table_styled, add_diagram_box, add_code_box,
    ROYAL_BLUE, NAVY, TEXT_DARK, SLATE
)
from viva_sections_part1 import add_six_step_block

def build_section_11(doc):
    add_heading_1(doc, "SECTION 11 — PHASE 5: HYBRID RETRIEVAL")
    
    add_six_step_block(doc,
        "Hybrid Retrieval (Score-Level Convex Fusion)",
        "Hybrid Retrieval is an Information Retrieval fusion strategy that combines normalized lexical scores (BM25) with dense semantic scores (Cosine Similarity) at the individual document level across the entire corpus.",
        "BM25 provides exact keyword precision and named-entity accuracy, while Semantic Retrieval provides conceptual synonym coverage. Hybrid retrieval combines both strengths, overcoming the vocabulary mismatch problem while retaining keyword precision.",
        "Because BM25 scores are unbounded positive numbers (e.g., 0 to 18) and Semantic scores are bounded in [0, 1], BM25 is first normalized using per-query Min-Max scaling. The normalized scores are then combined via convex linear combination: HybridScore = α * BM25_norm + (1 - α) * SemanticScore.",
        "Implemented in backend/retrieval/hybrid.py. Operates across the full 10,000-event corpus before selecting Top-K candidates. Hyperparameter α = 0.70 was empirically selected via a multi-alpha grid search.",
        "Query: 'machine learning workshops for students in Mumbai'. An event that matches the exact words 'machine learning' gets a high BM25 score. Another event titled 'Introductory Deep Learning Boot Camp' gets a high semantic score. The hybrid fusion model ranks both events highly, outperforming either single model.",
        "\"Hybrid retrieval is our dual-paradigm fusion engine. We solved score incommensurability by applying per-query Min-Max normalization to BM25 before fusing it with semantic cosine similarity. At α = 0.70, hybrid retrieval achieved our project's highest Recall@10 (0.2268) and highest F1 Score (0.2722), proving that combining lexical and neural models expands coverage without sacrificing precision.\""
    )

    add_heading_2(doc, "11.1 The Score Alignment & Incommensurability Problem")
    add_p(doc, "BM25 produces unbounded positive scores (typically ranging from 0.0 to 18.0+ depending on query length and token rarity). In contrast, Semantic Retrieval produces cosine similarity scores bounded strictly in [-1.0, 1.0] (practically [0.0, 0.90] for English text). If you simply add BM25 + Semantic, BM25 completely overwhelms the neural vector score! To combine them meaningfully, BM25 scores must be normalized on a per-query basis into [0.0, 1.0].")

    add_formula_box(doc, "Per-Query Min-Max BM25 Normalization & Convex Fusion",
        "BM25_norm(d) = [ BM25(d) - min_j BM25(j) ] / [ max_j BM25(j) - min_j BM25(j) ]\n"
        "HybridScore(d) = α * BM25_norm(d) + (1 - α) * SemanticScore(d)",
        [
            ("BM25(d)", "Raw BM25 score for document d"),
            ("min_j, max_j", "Minimum and maximum BM25 scores across all 10,000 documents for this specific query"),
            ("SemanticScore(d)", "Cosine similarity between query embedding and document embedding in [0, 1]"),
            ("α (Alpha = 0.70)", "Convex fusion hyperparameter balancing lexical precision (70%) and semantic recall (30%)"),
            ("Edge Case Handling", "If max_j == min_j (e.g. no terms match), BM25_norm is safely assigned 0.0 without division by zero")
        ]
    )

    add_code_box(doc, "Hybrid Fusion Implementation (backend/retrieval/hybrid.py)",
        "class HybridRetriever:\n"
        "    def __init__(self, bm25_retriever, semantic_retriever):\n"
        "        self.bm25 = bm25_retriever\n"
        "        self.semantic = semantic_retriever\n\n"
        "    def search(self, query, top_k=10, alpha=0.70):\n"
        "        # 1. Compute full-corpus scores for all 10,000 items\n"
        "        bm25_raw = self.bm25.get_all_scores(query)\n"
        "        sem_scores = self.semantic.get_all_scores(query)\n\n"
        "        # 2. Per-query Min-Max normalization\n"
        "        min_val, max_val = np.min(bm25_raw), np.max(bm25_raw)\n"
        "        if max_val > min_val:\n"
        "            bm25_norm = (bm25_raw - min_val) / (max_val - min_val)\n"
        "        else:\n"
        "            bm25_norm = np.zeros_like(bm25_raw)\n\n"
        "        # 3. Full-corpus convex combination\n"
        "        hybrid_scores = alpha * bm25_norm + (1.0 - alpha) * sem_scores\n"
        "        top_indices = np.argsort(hybrid_scores)[::-1][:top_k]\n"
        "        return top_indices, hybrid_scores[top_indices]"
    )

    add_heading_2(doc, "11.2 Empirical Hyperparameter Grid Search (Alpha Evaluation)")
    add_p(doc, "In Phase 5, an extensive empirical grid search was executed across all 500 query topics to determine the optimal alpha weight:")
    
    headers = ["Alpha (α)", "BM25 Weight", "Semantic Weight", "P@5", "P@10", "Recall@10", "F1", "MRR", "NDCG@10", "Analysis"]
    rows = [
        ["0.25", "25%", "75%", "0.3276", "0.2662", "0.1775", "0.2130", "0.4640", "0.2949", "Semantic heavy; poor keyword precision"],
        ["0.50", "50%", "50%", "0.3740", "0.3280", "0.2187", "0.2624", "0.4907", "0.3578", "Balanced; competitive across all metrics"],
        ["0.60", "60%", "40%", "0.3768", "0.3382", "0.2255", "0.2706", "0.4938", "0.3662", "High precision, strong semantic assistance"],
        ["0.70", "70%", "30%", "0.3808", "0.3402", "0.2268", "0.2722", "0.4915", "0.3689", "SELECTED: Highest P@10 (0.3402), Recall@10 (0.2268) & F1 (0.2722)"],
        ["0.75", "75%", "25%", "0.3820", "0.3392", "0.2261", "0.2714", "0.4888", "0.3686", "Lexical heavy; recall begins to decay"]
    ]
    add_table_styled(doc, headers, rows, [0.6, 0.7, 0.7, 0.6, 0.6, 0.6, 0.6, 0.6, 0.6, 1.4])

def build_section_12(doc):
    add_heading_1(doc, "SECTION 12 — PHASE 6: STRUCTURED EVENT RANKING")
    
    add_six_step_block(doc,
        "Structured Intent-Aware Ranking",
        "Structured Event Ranking is an explainable re-ranking layer that extracts structured constraints (city, category, admission, event type, audience) directly from natural-language queries using corpus vocabulary and applies an intent satisfaction boost.",
        "Even with Hybrid Retrieval, pure text matching cannot distinguish between an event taking place in Mumbai versus one mentioning Mumbai in a comparison. Furthermore, constraints like 'free entry' or 'for students' are hard constraints that users explicitly care about.",
        "The system scans the query against distinct entity values in the database (cities, categories, keywords like 'free'). Active signals are identified and assigned predefined weights. A normalized StructuredScore in [0, 1] is calculated, and blended with the Hybrid score at λ = 0.15.",
        "Implemented in backend/retrieval/structured_ranking.py. Dynamically checks query tokens against PostgreSQL vocabulary. Normalizes strictly by active signal weights so unmentioned signals don't penalize candidate events.",
        "Query: 'free technology workshops in Mumbai'. Active signals: City='Mumbai' (0.30), Category='Technology' (0.25), Free=True (0.20), EventType='Workshop' (0.15). An event meeting all 4 criteria receives StructuredScore = 1.0000, boosting its final rank.",
        "\"Structured Ranking bridges natural language search with structured database constraints. Instead of brittle regex rules, it matches tokens against the database vocabulary and applies a soft additive boost at lambda=0.15. This pushed our top-rank quality metrics to their project peaks: MRR reached 0.4991 and NDCG@10 reached 0.3731.\""
    )

    add_heading_2(doc, "12.1 Rule-Free, Vocabulary-Driven Intent Extraction")
    add_p(doc, "EventIQ dynamically extracts query intent by checking candidate tokens against the actual vocabulary of the PostgreSQL database:")
    add_bullet(doc, "Matches tokens against all 20 distinct cities in events.city. Multi-word cities are checked before single-word prefixes (e.g. 'New Delhi' before 'Delhi').", bold_prefix="1. City Signal (Weight = 0.30): ")
    add_bullet(doc, "Matches tokens against the 20 distinct categories in events.category (e.g. 'technology', 'photography', 'gaming').", bold_prefix="2. Category Signal (Weight = 0.25): ")
    add_bullet(doc, "Matches queries containing 'free', 'no cost', 'complimentary' -> checks event.is_free == True.", bold_prefix="3. Admission Signal (Weight = 0.20): ")
    add_bullet(doc, "Matches tokens against singular/plural event types ('workshop/workshops', 'conference/conferences', 'meetup/meetups').", bold_prefix="4. Event Type Signal (Weight = 0.15): ")
    add_bullet(doc, "Matches terms like 'students', 'beginners', 'developers', 'researchers', 'professionals'.", bold_prefix="5. Audience Signal (Weight = 0.10): ")

    add_formula_box(doc, "Structured Score Formula",
        "StructuredScore(q, d) = [ sum_{s in ActiveSignals(q)} weight(s) * match(s, d) ] / [ sum_{s in ActiveSignals(q)} weight(s) ]",
        [
            ("ActiveSignals(q)", "Set of intent signals detected in query q"),
            ("weight(s)", "Predefined initial weight: City=0.30, Category=0.25, Free=0.20, Type=0.15, Audience=0.10"),
            ("match(s, d)", "Binary match indicator (1 if event d satisfies signal s, 0 otherwise)"),
            ("Denominator", "Sum of weights of active signals; ensures normalization to [0.0, 1.0]")
        ]
    )

    add_code_box(doc, "Structured Ranker Intent Extraction (backend/retrieval/structured_ranking.py)",
        "class StructuredEventRanker:\n"
        "    SIGNAL_WEIGHTS = {\n"
        "        'city': 0.30, 'category': 0.25, 'is_free': 0.20,\n"
        "        'event_type': 0.15, 'audience': 0.10\n"
        "    }\n\n"
        "    def rank(self, query, top_k=10, lambda_weight=0.15):\n"
        "        # 1. Retrieve candidates via HybridRetriever (alpha=0.70)\n"
        "        hybrid_indices, hybrid_scores = self.hybrid.search(query, top_k=1000)\n"
        "        # 2. Extract active query intents from database vocabulary\n"
        "        active_intents = self.extract_intent(query)\n"
        "        # 3. Compute StructuredScore & Blend with HybridScore (lambda=0.15)\n"
        "        # FinalScore = (1 - lambda_weight) * HybridScore + lambda_weight * StructuredScore\n"
        "        return final_ranked_results[:top_k]"
    )

def build_section_13(doc):
    add_heading_1(doc, "SECTION 13 — THE FINAL RANKING FORMULA")
    
    add_formula_box(doc, "EventIQ Final Ranking Formula",
        "FinalScore(d) = (1 - λ) * HybridScore(d) + λ * StructuredScore(d)\n"
        "FinalScore(d) = 0.85 * HybridScore(d) + 0.15 * StructuredScore(d)",
        [
            ("0.85 (85%)", "Weight assigned to the Hybrid Retrieval score (lexical + semantic relevance)"),
            ("0.15 (15%)", "Weight assigned to the Structured Intent satisfaction score"),
            ("λ = 0.15", "Re-ranking hyperparameter selected to promote exact constraint matches to top rank positions without disrupting core topical relevance")
        ]
    )

    add_heading_2(doc, "13.1 Step-by-Step Numerical Calculation Example")
    add_p(doc, "Suppose a user searches for: 'free photography workshops for beginners in Mumbai'.")
    add_p(doc, "Event candidate: EVT008120 (Title: 'Beginner Portrait Workshop', City: 'Mumbai', Category: 'Photography', is_free: True, EventType: 'Workshop', Audience: 'Beginners').")
    add_bullet(doc, "Suppose raw normalized BM25 = 0.8800, Semantic Cosine = 0.8200 -> HybridScore = 0.70(0.8800) + 0.30(0.8200) = 0.6160 + 0.2460 = 0.8620.", bold_prefix="Step 1: Compute HybridScore: ")
    add_bullet(doc, "Active signals detected: City (0.30), Category (0.25), Free (0.20), EventType (0.15), Audience (0.10). Sum of active weights = 1.00.", bold_prefix="Step 2: Identify Active Intent Signals: ")
    add_bullet(doc, "The event matches all 5 signals -> StructuredScore = (0.30*1 + 0.25*1 + 0.20*1 + 0.15*1 + 0.10*1) / 1.00 = 1.0000.", bold_prefix="Step 3: Compute StructuredScore: ")
    add_bullet(doc, "FinalScore = 0.85 * (0.8620) + 0.15 * (1.0000) = 0.7327 + 0.1500 = 0.8827.", bold_prefix="Step 4: Compute FinalScore: ")
    add_bullet(doc, "The structured layer boosts this fully compliant event from 0.8620 to 0.8827, placing it at Rank #1 above events that match the keywords but have admission fees or are situated in other cities.", bold_prefix="Result: ")

def build_section_14(doc):
    add_heading_1(doc, "SECTION 14 — EXPLAINABILITY & SCORE DECOMPOSITION")
    add_p(doc, "Unlike black-box recommendation systems, EventIQ provides complete mathematical explainability. For every retrieved result, the user can inspect the exact score decomposition:")
    
    add_heading_2(doc, "14.1 Returned Explainability Metadata")
    add_bullet(doc, "The blended rank score used to position the card.", bold_prefix="1. Final Score: ")
    add_bullet(doc, "The 85% foundation score derived from BM25 and semantic fusion.", bold_prefix="2. Hybrid Score: ")
    add_bullet(doc, "The 15% intent compliance score.", bold_prefix="3. Structured Score: ")
    add_bullet(doc, "Key-value dictionary showing matched constraints (e.g. {'city': 'Mumbai', 'category': 'Photography', 'is_free': True}).", bold_prefix="4. Matched Signals: ")
    add_bullet(doc, "Exact additive point contributions for each factor (e.g. hybrid_base: 0.7327, city_match: 0.0450, is_free_match: 0.0300).", bold_prefix="5. Signal Contributions: ")

    add_heading_2(doc, "14.2 Frontend Visualization in React")
    add_p(doc, "In the React frontend, clicking 'Explain Score' on any card opens an inline glassmorphic panel displaying the mathematical equation, green checkmark badges for active matched intents, and animated horizontal progress bars showing the proportional contribution of each signal to the final score.")

def build_section_15(doc):
    add_heading_1(doc, "SECTION 15 — COMPLETE END-TO-END SEARCH FLOW")
    add_p(doc, "The following sequence diagram illustrates the complete execution path when a user submits a search query in EventIQ:")

    diagram = (
        "USER                 REACT FRONTEND              FASTAPI BACKEND          SEARCHSERVICE         DATABASE / EMBEDDINGS\n"
        " │                         │                            │                       │                        │\n"
        " ├── Types Query ─────────>│                            │                       │                        │\n"
        " │   & Clicks 'Retrieve'   ├── POST /api/search ───────>│                       │                        │\n"
        " │                         │   {query, model, top_k}    ├── Validate Schema     │                        │\n"
        " │                         │                            │   via Pydantic        │                        │\n"
        " │                         │                            ├── search() ──────────>│                        │\n"
        " │                         │                            │                       ├── 1. Compute BM25      │\n"
        " │                         │                            │                       │      over 10k docs     │\n"
        " │                         │                            │                       ├── 2. Encode Query &    │\n"
        " │                         │                            │                       │      Cosine Similarity ├── (10k x 384 npy)\n"
        " │                         │                            │                       ├── 3. Min-Max Normalize │\n"
        " │                         │                            │                       │      & Fuse (alpha=.7) │\n"
        " │                         │                            │                       ├── 4. Extract Intent    │\n"
        " │                         │                            │                       │      & Rank (lambda=.15)│\n"
        " │                         │                            │                       ├── 5. Extract Top-K &   │\n"
        " │                         │                            │                       │      Fetch Metadata ──>├── (events table)\n"
        " │                         │                            │<── Ranked Results ────┤                        │\n"
        " │                         │<── JSON Response (HTTP 200)┤                       │                        │\n"
        " │                         │    {results, scores, etc.} │                       │                        │\n"
        " │                         ├── Render EventCards        │                       │                        │\n"
        " │<── Displays Results ────┤   & Compute Facet Counts   │                       │                        │\n"
        " │                         │                            │                       │                        │"
    )
    add_diagram_box(doc, "Complete Query Processing Sequence Diagram", diagram)

def build_section_16(doc):
    add_heading_1(doc, "SECTION 16 — FASTAPI BACKEND ARCHITECTURE")
    add_p(doc, "The EventIQ backend is implemented in Python using FastAPI, organized into modular packages:")
    add_bullet(doc, "ASGI application definition, CORS middleware, global exception handlers, and route handlers.", bold_prefix="backend/app/main.py: ")
    add_bullet(doc, "Pydantic models defining input validation constraints and output serialization schemas.", bold_prefix="backend/app/schemas.py: ")
    add_bullet(doc, "PostgreSQL connection pool management and query execution helper functions.", bold_prefix="backend/app/database.py: ")
    add_bullet(doc, "In-memory singleton holding loaded retrieval models and orchestrating search execution.", bold_prefix="backend/app/services/search_service.py: ")
    add_bullet(doc, "Implementations of TF-IDF, BM25, Semantic, Hybrid, and Structured ranking algorithms.", bold_prefix="backend/retrieval/: ")

    add_heading_2(doc, "16.1 REST Endpoints Specification")
    headers = ["Endpoint", "HTTP Method", "Request Payload", "Response Summary", "Description"]
    rows = [
        ["/api/health", "GET", "None", "{\"status\": \"ok\", \"database\": \"connected\"}", "Service liveness and database connectivity verification."],
        ["/api/events/count", "GET", "None", "{\"event_count\": 10000}", "Returns total number of events indexed in PostgreSQL."],
        ["/api/categories", "GET", "None", "[{\"category\": \"Technology\", \"count\": 461}, ...]", "Returns 20 distinct event categories with occurrence counts."],
        ["/api/locations", "GET", "None", "[{\"city\": \"Mumbai\", \"state\": \"...\", \"count\": 508}, ...]", "Returns 20 distinct cities with state and occurrence counts."],
        ["/api/evaluation", "GET", "None", "{\"comparison\": [...], \"raw_benchmarks\": {...}}", "Serves reproducible benchmark metrics across all five models."],
        ["/api/events/{id}", "GET", "Path: event_id", "{\"event_id\": \"EVT001305\", ... (22 fields)}", "Retrieves complete metadata schema for an individual event."],
        ["/api/search", "POST", "{\"query\": \"...\", \"model\": \"structured\", \"top_k\": 10}", "{\"results\": [...], \"total_results\": 10, ...}", "Executes multi-paradigm event retrieval and ranking."]
    ]
    add_table_styled(doc, headers, rows, [1.1, 0.7, 1.4, 1.7, 1.9])

    add_code_box(doc, "Search Service Singleton (backend/app/services/search_service.py)",
        "class SearchService:\n"
        "    _instance = None\n\n"
        "    @classmethod\n"
        "    def get_instance(cls):\n"
        "        if cls._instance is None:\n"
        "            cls._instance = cls()\n"
        "        return cls._instance\n\n"
        "    def __init__(self):\n"
        "        # Preload TF-IDF, BM25, Semantic Embeddings, Hybrid & Structured models once\n"
        "        self.tfidf_retriever = TFIDFRetriever(corpus)\n"
        "        self.bm25_retriever = BM25Retriever(tokenized_corpus)\n"
        "        self.semantic_retriever = SemanticRetriever(embeddings_path)\n"
        "        self.hybrid_retriever = HybridRetriever(self.bm25_retriever, self.semantic_retriever)\n"
        "        self.structured_ranker = StructuredEventRanker(self.hybrid_retriever, events_df)"
    )

def build_section_17(doc):
    add_heading_1(doc, "SECTION 17 — REACT FRONTEND ARCHITECTURE")
    add_p(doc, "The frontend is built with React 19, Vite, and TypeScript. It strictly adheres to separation of concerns: zero retrieval or ranking algorithms are implemented in TypeScript. All scoring decisions occur on the Python backend; the frontend functions strictly as an interactive consumer and visualizer.")

    add_heading_2(doc, "17.1 Component Hierarchy")
    add_bullet(doc, "Root layout manager controlling active view navigation ('search', 'evaluation', 'architecture'), live API health polling, and event detail modal state.", bold_prefix="App.tsx: ")
    add_bullet(doc, "Navigation header bar displaying EventIQ branding, view tabs, live API connection dot, and 10,000-event badge.", bold_prefix="Header.tsx: ")
    add_bullet(doc, "Main search view coordinating search input, model selection, faceted sidebar filtering, and card list rendering.", bold_prefix="SearchView.tsx: ")
    add_bullet(doc, "Search input with clear button, submit button, and 6 curated one-click query chips.", bold_prefix="SearchBar.tsx: ")
    add_bullet(doc, "Pill buttons to toggle between Structured, Hybrid, BM25, Semantic, and TF-IDF engines, plus Top-K selector.", bold_prefix="ModelSelector.tsx: ")
    add_bullet(doc, "Dynamically aggregates categories, cities, and pricing across current retrieved candidates.", bold_prefix="FiltersSidebar.tsx: ")
    add_bullet(doc, "Individual event listing displaying rank, score badge, title, category tags, price/free chip, and action buttons.", bold_prefix="EventCard.tsx: ")
    add_bullet(doc, "Inline expandable panel revealing mathematical formulas, active matched intent chips, and signal contribution progress bars.", bold_prefix="ExplainabilityPanel.tsx: ")
    add_bullet(doc, "Comprehensive modal window querying GET /api/events/{id} to display the full 22-attribute record.", bold_prefix="EventDetailsModal.tsx: ")
    add_bullet(doc, "Interactive evaluation dashboard rendering the academic metrics table and Recharts bar graphs.", bold_prefix="EvaluationView.tsx: ")
    add_bullet(doc, "System documentation view explaining dataset attributes, retrieval models, and mathematical formulas.", bold_prefix="ArchitectureView.tsx: ")
    add_bullet(doc, "Centralized API client consuming FastAPI endpoints via fetch() with error interception.", bold_prefix="services/api.ts: ")

def build_section_18(doc):
    add_heading_1(doc, "SECTION 18 — THE SEARCH INTERFACE & FILTERING")
    add_p(doc, "The Search page (SearchView.tsx) is designed for intuitive exploration. Key interactive elements include:")
    add_bullet(doc, "Pre-populated natural-language topics (e.g. 'free machine learning workshops in Mumbai', 'music festivals in Goa') that allow users to test the pipeline with one click.", bold_prefix="1. Example Query Chips: ")
    add_bullet(doc, "Allows instant switching between all 5 retrieval algorithms without re-typing the query, enabling immediate side-by-side ranking comparisons.", bold_prefix="2. Multi-Model Selector: ")
    add_bullet(doc, "Configurable from 5 to 50 results (default: 10).", bold_prefix="3. Top-K Control: ")
    add_bullet(doc, "Users can sort the retrieved results by Rank/Score, Rating (High to Low), Price (Low to High), or Date.", bold_prefix="4. Dynamic Sorting: ")

    add_callout(doc, "CRITICAL EXPLANATION: How Facet Filtering Works in EventIQ",
        "Examiners may ask: 'Do your sidebar category and city filters change the ranking algorithm?'\n\n"
        "Your precise answer:\n"
        "\"No. In EventIQ, the sidebar filters operate strictly over the retrieved candidate set. When the backend returns the Top-K events for a query, the frontend dynamically aggregates the categories, cities, and admission types present in that set. If the user clicks 'Free' or 'Technology', the UI filters the displayed cards among those retrieved candidates. This preserves the academic purity of the underlying IR models while providing users with intuitive client-side exploration.\"",
        "teacher"
    )

def build_section_19(doc):
    add_heading_1(doc, "SECTION 19 — EVENT DETAILS MODAL")
    add_p(doc, "When a user clicks 'View Details' on an event card, the frontend issues a GET /api/events/{event_id} request to fetch the complete un-truncated entity record from PostgreSQL. The modal presents a comprehensive specifications table covering all 22 attributes, including coordinates, organizer name, language, expected attendees, venue address, and rating.")

def build_section_20(doc):
    add_heading_1(doc, "SECTION 20 — EVALUATION DASHBOARD & METRICS THEORY")
    add_p(doc, "Evaluation is the scientific core of Information Retrieval. EventIQ evaluated all five retrieval paradigms over 500 standardized query topics and 7,500 relevance judgments using six industry-standard TREC metrics.")

    add_heading_2(doc, "20.1 Detailed Metrics Breakdown")
    metrics_info = [
        ("Precision@5 (P@5)", "P@5 = (Relevant items in top 5) / 5", "Fraction of retrieved events in the first 5 slots that are relevant. Measures immediate user satisfaction.", "BM25 leads at 0.3860."),
        ("Precision@10 (P@10)", "P@10 = (Relevant items in top 10) / 10", "Fraction of retrieved events in the top 10 slots that are relevant. Measures first-page precision.", "Hybrid leads at 0.3402."),
        ("Recall@10 (R@10)", "R@10 = (Relevant items in top 10) / (Total relevant items in corpus)", "Fraction of all relevant events in the corpus that were successfully retrieved in the top 10 slots.", "Hybrid leads at 0.2268."),
        ("F1 Score (Top-10)", "F1 = 2 * (P@10 * R@10) / (P@10 + R@10)", "Harmonic mean of Precision@10 and Recall@10. Balances precision and recall into a single score.", "Hybrid leads at 0.2722."),
        ("Mean Reciprocal Rank (MRR)", "MRR = (1 / |Q|) * sum_{q} (1 / rank_first_relevant)", "Reciprocal rank of the very first relevant document retrieved. Measures how quickly the user finds their first good match.", "Structured leads at 0.4991."),
        ("NDCG@10", "NDCG@10 = DCG@10 / IDCG@10, where DCG@10 = sum_{i=1}^{10} (2^{rel_i} - 1) / log2(i + 1)", "Normalized Discounted Cumulative Gain. Considers graded relevance (0-3) and penalizes relevant items that appear at lower ranks.", "Structured leads at 0.3731.")
    ]
    for name, form, desc, leader in metrics_info:
        add_p(doc, f"{desc} Formula: {form} ({leader})", bold_prefix=f"{name}: ")

    add_heading_2(doc, "20.2 Comparative Benchmark Performance Summary")
    headers = ["Model Architecture", "P@5", "P@10", "Recall@10", "F1 Score", "MRR", "NDCG@10", "Primary Strength"]
    rows = [
        ["TF-IDF Baseline", "0.3552", "0.2628", "0.1752", "0.2102", "0.4820", "0.3101", "Standard lexical benchmark"],
        ["BM25 Lexical", "0.3860", "0.3250", "0.2167", "0.2600", "0.4797", "0.3623", "Highest Precision@5 (0.3860)"],
        ["Semantic (all-MiniLM-L6-v2)", "0.2248", "0.1788", "0.1192", "0.1430", "0.3729", "0.1955", "Captures conceptual associations"],
        ["Hybrid (α = 0.70)", "0.3808", "0.3402", "0.2268", "0.2722", "0.4915", "0.3689", "Highest P@10 (0.3402), R@10 (0.2268) & F1 (0.2722)"],
        ["Structured Ranking (λ = 0.15)", "0.3832", "0.3364", "0.2243", "0.2691", "0.4991", "0.3731", "Highest MRR (0.4991) & NDCG@10 (0.3731)"]
    ]
    add_table_styled(doc, headers, rows, [1.4, 0.6, 0.6, 0.7, 0.7, 0.6, 0.7, 1.5])

